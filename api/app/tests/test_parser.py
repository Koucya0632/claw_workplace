from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.parsers.file_parser import FileParser


def build_docx_bytes() -> bytes:
    # 以記憶體生成 docx，避免測試依賴外部檔案。
    document = Document()
    document.add_paragraph("OpenClaw 會議摘要")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx_bytes() -> bytes:
    # 以記憶體建立最小 xlsx，驗證 parser 能讀出工作表內容。
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "summary"
    worksheet.append(["month", "revenue"])
    worksheet.append(["2026-03", 133000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_pdf_bytes() -> bytes:
    # PDF 測試先用最小空白頁，重點是確保 parser 能處理合法 PDF。
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("path", "payload", "expected_extension"),
    [
        (Path("note.txt"), "OpenClaw txt".encode("utf-8"), "txt"),
        (Path("note.md"), "# OpenClaw md".encode("utf-8"), "md"),
        (Path("table.csv"), "name,value\nrevenue,100".encode("utf-8"), "csv"),
        (Path("table.xlsx"), build_xlsx_bytes(), "xlsx"),
        (Path("memo.docx"), build_docx_bytes(), "docx"),
        (Path("paper.pdf"), build_pdf_bytes(), "pdf"),
    ],
)
def test_file_parser_supports_phase_one_formats(path: Path, payload: bytes, expected_extension: str) -> None:
    # 每一種 Phase 1 支援格式都至少驗證一次成功解析。
    parser = FileParser()
    parsed = parser.parse(path, payload)

    assert parsed.extension == expected_extension
    assert parsed.mime_type


def test_file_parser_raises_for_corrupted_docx() -> None:
    # 損壞檔案必須要明確丟出例外，索引服務才知道要記錄錯誤。
    parser = FileParser()

    with pytest.raises(Exception):  # noqa: B017
        parser.parse(Path("broken.docx"), b"not-a-valid-docx")

